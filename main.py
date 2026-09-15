import flet as ft
import os, sys, time, re, csv, hashlib, shutil, urllib.request, tempfile, glob
import xml.etree.ElementTree as ET
import openpyxl
from bs4 import BeautifulSoup
import socks
from sockshandler import SocksiPyHandler
from datetime import datetime, timedelta

# --- Constants & Regex ---
CURRENT_YEAR = datetime.now().year
RE_USP, RE_EP, RE_BP = re.compile(r'\bUSP\b', re.I), re.compile(r'\bEP\b', re.I), re.compile(r'\bBP\b', re.I)
EXTENDED_STEMS = ['MONOHYDRATE', 'DIHYDRATE', 'HYDROCHLORIDE', 'HCL', 'SODIUM', 'POTASSIUM', 'FOR ASSAY', 'CRS', 'BRP', 'RS']

# --- Core Logic Functions ---
def sanitize_batch(raw_val: str) -> str:
    s = str(raw_val).strip().upper()
    if s.endswith(".0"): s = s[:-2]
    s = re.sub(r'^(LOT|BATCH|B\.?\s*NO\.?|#)\s*[:.-]?\s*', '', s)
    if s.isdigit(): s = str(int(s))
    return s

def safe_normalize(name):
    return re.sub(r'[^A-Z0-9%]', '', str(name).upper().replace('\n', ' ').strip())

def extract_targets_from_stock(filepath, mode, skip_dated=False):
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    raw_targets = []
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        name_col, batch_col, date_col, code_col = -1, -1, -1, -1
        for row in ws.iter_rows(values_only=True):
            if name_col == -1:
                for c_idx, cell in enumerate(row):
                    if cell and isinstance(cell, str):
                        c_low = cell.lower()
                        if "name of reference" in c_low: name_col = c_idx
                        elif "batch/lot" in c_low: batch_col = c_idx
                        elif "pharmacopeial" in c_low: date_col = c_idx
                        elif "code" in c_low: code_col = c_idx
                if name_col != -1: continue 
            
            if name_col != -1 and batch_col != -1:
                name_val = str(row[name_col]).strip() if row[name_col] else ""
                batch_val = sanitize_batch(row[batch_col]) if row[batch_col] is not None else ""
                code_val = str(row[code_col]).strip() if code_col != -1 and row[code_col] else ""
                
                if name_val and name_val.lower() != "nan" and "name of reference" not in name_val.lower():
                    if mode == "1" and RE_USP.search(name_val) and batch_val: raw_targets.append((batch_val, name_val, code_val))
                    elif mode == "2" and RE_EP.search(name_val): raw_targets.append((batch_val, name_val, code_val)) 
                    elif mode == "3" and RE_BP.search(name_val) and batch_val: raw_targets.append((batch_val, name_val, code_val))
    
    unique_targets = {}
    for t_val, t_name, t_code in raw_targets:
        key = f"{t_name.strip().upper()}_{t_val.strip().upper()}"
        if key not in unique_targets: unique_targets[key] = (t_val, t_name, t_code)
    return list(unique_targets.values())

def ping_network(route):
    try:
        start = time.time()
        if route == "ipvanish":
            opener = urllib.request.build_opener(SocksiPyHandler(socks.SOCKS5, "nyc.socks.ipvanish.com", 1080, True, "Snyc9enhuitq", "GVMTGhFxz"))
            urllib.request.install_opener(opener)
        else:
            urllib.request.install_opener(urllib.request.build_opener())
        
        req = urllib.request.Request("https://www.pharmacopoeia.com", headers={'User-Agent': 'Mozilla/5.0'})
        urllib.request.urlopen(req, timeout=5)
        return (time.time() - start) * 1000
    except:
        return float('inf')

# --- Flet UI Application ---
def main(page: ft.Page):
    page.title = "Omni Auditor"
    page.theme_mode = ft.ThemeMode.DARK
    page.padding = 20
    page.window_width = 400
    page.window_height = 800

    app_state = {
        "mother_excel": None,
        "catalogs": {"USP": False, "EP": False, "BP": False},
        "sources": {},
        "network": None
    }

    # --- Pickers ---
    def on_mother_picked(e: ft.FilePickerResultEvent):
        if e.files:
            app_state["mother_excel"] = e.files[0].path
            mother_btn.text = f"Loaded: {e.files[0].name}"
            mother_btn.icon = ft.icons.CHECK_CIRCLE
            mother_btn.icon_color = "green"
            page.update()

    mother_picker = ft.FilePicker(on_result=on_mother_picked)
    page.overlay.append(mother_picker)

    # --- Page 1: Setup ---
    mother_btn = ft.ElevatedButton("Upload Mother Excel", icon=ft.icons.UPLOAD_FILE, width=300, height=60, on_click=lambda _: mother_picker.pick_files(allowed_extensions=["xlsx", "xls"]))
    usp_chk = ft.Checkbox(label="USP Catalog", value=False)
    ep_chk = ft.Checkbox(label="EP Catalog", value=False)
    bp_chk = ft.Checkbox(label="BP Catalog", value=False)

    def go_source_routing(e):
        app_state["catalogs"] = {"USP": usp_chk.value, "EP": ep_chk.value, "BP": bp_chk.value}
        if not app_state["mother_excel"] or not any(app_state["catalogs"].values()): return
        build_source_page()
        page.go("/sources")

    # --- Page 2: Sources ---
    source_col = ft.Column(spacing=20)
    def build_source_page():
        source_col.controls.clear()
        for cat, checked in app_state["catalogs"].items():
            if checked:
                rg = ft.RadioGroup(content=ft.Row([ft.Radio(value="online", label="Check Online"), ft.Radio(value="local", label="Upload File")]), value="online")
                app_state["sources"][cat] = rg
                source_col.controls.append(ft.Card(content=ft.Container(padding=15, content=ft.Column([ft.Text(f"{cat} Catalog Source", weight="bold"), rg]))))
        page.update()

    # --- Page 3: Network Wizard ---
    net_rg = ft.RadioGroup(
        content=ft.Column([
            ft.Radio(value="ipvanish", label="IPVanish VPN"),
            ft.Radio(value="normal", label="Normal Connection"),
            ft.Radio(value="cloudflare", label="Cloudflare DNS"),
            ft.Radio(value="fastest", label="Pick Fastest (Auto-Ping)"),
        ]), value="fastest"
    )

    log_view = ft.ListView(expand=True, spacing=10, height=300)
    def log(msg, color="white"):
        log_view.controls.append(ft.Text(msg, color=color, size=12))
        log_view.scroll_to(offset=-1, duration=100)
        page.update()

    def execute_audit(e):
        app_state["network"] = net_rg.value
        page.go("/executing")
        
        # 1. Network Test
        selected_net = app_state["network"]
        if selected_net == "fastest":
            log("📡 Auto-pinging network routes...")
            routes = {"normal": ping_network("normal"), "ipvanish": ping_network("ipvanish")}
            best = min(routes, key=routes.get)
            log(f"⚡ Fastest route locked: {best.upper()} ({int(routes[best])}ms)", "green")
            selected_net = best
            
        if selected_net == "ipvanish":
            opener = urllib.request.build_opener(SocksiPyHandler(socks.SOCKS5, "nyc.socks.ipvanish.com", 1080, True, "Snyc9enhuitq", "GVMTGhFxz"))
            urllib.request.install_opener(opener)
        else:
            urllib.request.install_opener(urllib.request.build_opener())

        # 2. Extract Data
        log("📊 Extracting stock records from Excel...")
        targets = {}
        modes = {"USP": "1", "EP": "2", "BP": "3"}
        for cat, checked in app_state["catalogs"].items():
            if checked:
                targets[cat] = extract_targets_from_stock(app_state["mother_excel"], modes[cat], skip_dated=True)
                log(f"✔ Extracted {len(targets[cat])} {cat} items.", "cyan")

        # 3. Generate Dashboard
        log("🎨 Building Interactive HTML Dashboard...", "yellow")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        download_dir = "/storage/emulated/0/Download" if os.path.exists("/storage/emulated/0/Download") else os.path.expanduser("~")
        html_path = os.path.join(download_dir, f"Omni_Audit_{timestamp}.html")
        
        html_content = f"""
        <!DOCTYPE html><html><head><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Omni Hub</title>
        <style>
            body {{ font-family: sans-serif; background: #0b0e14; color: #e2e8f0; padding: 15px; margin: 0; }}
            .app-bar {{ background: #151922; padding: 20px; border-radius: 20px; border: 1px solid #2d3548; margin-bottom: 20px; }}
            .title {{ font-size: 1.5rem; font-weight: bold; margin: 0; color: #fff; }}
            .subtitle {{ font-size: 0.85rem; color: #00e5a3; margin-top: 5px; }}
            .card {{ background: #151922; padding: 15px; border-radius: 15px; border: 1px solid #2d3548; margin-bottom: 10px; }}
            .stat-row {{ display: flex; justify-content: space-between; font-size: 0.85rem; color: #94a3b8; margin-top: 8px; padding-top: 8px; border-top: 1px solid #2d3548; }}
        </style></head><body>
        <div class="app-bar">
            <h1 class="title">Omni Auditor Hub</h1>
            <div class="subtitle">Beximco Pharmaceuticals PLC • QC Dashboard (Tongi Site)</div>
            <div style="font-size: 0.75rem; color: #8a99ad; margin-top: 5px;">Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}</div>
        </div>
        """
        for cat in targets:
            html_content += f"<h2 style='color:#4e8cff;'>{cat} Extraction</h2>"
            for t_val, t_name, t_code in targets[cat]:
                html_content += f"<div class='card'><b>{t_name}</b><div class='stat-row'><span>Code: {t_code}</span><span>Batch: {t_val}</span></div></div>"
        html_content += "</body></html>"
        
        try:
            with open(html_path, "w", encoding="utf-8") as f: f.write(html_content)
            log(f"💾 Saved successfully to Downloads:\n{html_path}", "green")
            page.launch_url(f"file://{html_path}")
        except Exception as e:
            log(f"❌ Could not save file: {e}", "red")

    # --- Routing ---
    def route_change(route):
        page.views.clear()
        page.views.append(ft.View("/", [
            ft.Text("Omni Auditor", size=30, weight="heavy", color=ft.colors.BLUE_400),
            mother_btn, ft.Divider(height=10, color="transparent"),
            ft.Text("Select Catalogs:", weight="bold"), usp_chk, ep_chk, bp_chk,
            ft.ElevatedButton("Next", on_click=go_source_routing, width=300, bgcolor=ft.colors.BLUE_700, color="white")
        ], horizontal_alignment=ft.CrossAxisAlignment.CENTER))

        if page.route == "/sources":
            page.views.append(ft.View("/sources", [
                ft.Text("Source Routing", size=24, weight="bold"), source_col,
                ft.ElevatedButton("Next", on_click=lambda _: page.go("/network"), width=300, bgcolor=ft.colors.BLUE_700, color="white")
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER))

        if page.route == "/network":
            page.views.append(ft.View("/network", [
                ft.Text("Network Logic", size=24, weight="bold"),
                ft.Card(content=ft.Container(padding=20, content=net_rg)),
                ft.ElevatedButton("Execute Audit", on_click=lambda e: page.run_task(execute_audit), width=300, bgcolor=ft.colors.GREEN_600, color="white")
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER))

        if page.route == "/executing":
            page.views.append(ft.View("/executing", [
                ft.Text("Processing Audit...", size=24, weight="bold"),
                ft.ProgressRing(), log_view
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER))
        page.update()

    def view_pop(view):
        page.views.pop()
        page.go(page.views[-1].route)

    page.on_route_change = route_change
    page.on_view_pop = view_pop
    page.go(page.route)

ft.app(target=main)
